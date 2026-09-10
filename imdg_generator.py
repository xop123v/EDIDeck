# imdg_generator.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

def add_header(ws, headers, fill_color="DDDDDD"):
    bold = Font(bold=True)
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor=fill_color)
        cell.alignment = Alignment(horizontal="center")

def generate_imdg_sheet(wb, containers):
    ws = wb.create_sheet("IMDG")
    headers = ["Cont. ID","PoL","PoD","UN No","Class","Net Weight kg","Packing group","Flash Pt","EmS","Proper shipping name","Remark"]
    add_header(ws, headers, fill_color="FFDDDD")
    row = 2
    for c in containers:
        if c.get("dg"):
            ws.append([
                c.get("container"), c.get("pol"), c.get("pod"), c.get("unno"), c.get("class"),
                c.get("weight"), c.get("packing_group"), c.get("flash_point"), c.get("ems"),
                c.get("psn"), c.get("remark")
            ])
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(50, max_len + 2)
