# reefer_generator.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

def add_header(ws, headers, fill_color="DDFFDD"):
    bold = Font(bold=True)
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor=fill_color)
        cell.alignment = Alignment(horizontal="center")

def generate_reefer_sheet(wb, containers):
    ws = wb.create_sheet("Reefer")
    headers = ["Cont. ID","Bay","Row","Tier","Size","Set Temp","Ventilation","Remark"]
    add_header(ws, headers, fill_color="DDFFDD")
    for c in containers:
        if c.get("reefer"):
            ws.append([
                c.get("container"), c.get("bay"), c.get("row"), c.get("tier"), c.get("size"),
                c.get("temperature"), c.get("ventilation"), c.get("remark")
            ])
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(40, max_len + 2)
